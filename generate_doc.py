import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_callout(doc, text, title="NOTA IMPORTANTE", border_color="005A9C", fill_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(0, 90, 156)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_image_placeholder(doc, title, subtitle="Haga clic aquí y presione Ctrl + V para pegar la imagen o use el menú 'Insertar > Imagen' en Word."):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=260, bottom=260, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="dashed" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    
    r_icon = p.add_run("[ 📷 ESPACIO RESERVADO PARA CAPTURA DE PANTALLA ]\n")
    r_icon.bold = True
    r_icon.font.name = "Calibri"
    r_icon.font.size = Pt(10.5)
    r_icon.font.color.rgb = RGBColor(0, 79, 159)
    
    r_title = p.add_run(f"Figura: {title}\n")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(11.5)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p.add_run(f"{subtitle}\n(Dimensión recomendada: Ancho 14 cm a 16 cm)")
    r_sub.italic = True
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def style_table(tbl, col_widths, headers, data, header_bg="004F9F", alt_bg="F8FAFC"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    if len(tbl.rows) == 0:
        hdr_row = tbl.add_row()
    else:
        hdr_row = tbl.rows[0]
    hdr_cells = hdr_row.cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    for row_idx, row_data in enumerate(data):
        row = tbl.add_row()
        bg = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            c = row.cells[col_idx]
            c.text = str(cell_value)
            c.width = Inches(col_widths[col_idx])
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            p = c.paragraphs[0]
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(51, 65, 85)
                
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_document(output_path):
    doc = Document()
    
    # Margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    
    # ================= COVER PAGE =================
    cover_p1 = doc.add_paragraph()
    cover_p1.paragraph_format.space_before = Pt(60)
    cover_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = cover_p1.add_run("NINABIT SOLUTIONS  |  SKYWORTH BOLIVIA")
    r_tag.bold = True
    r_tag.font.size = Pt(12)
    r_tag.font.color.rgb = RGBColor(0, 163, 224) # Cyan
    
    cover_p2 = doc.add_paragraph()
    cover_p2.paragraph_format.space_before = Pt(14)
    cover_p2.paragraph_format.space_after = Pt(14)
    cover_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = cover_p2.add_run("DOCUMENTACIÓN TÉCNICA, ARQUITECTURA Y MANUAL DE MÓDULOS")
    r_title.bold = True
    r_title.font.size = Pt(23)
    r_title.font.color.rgb = RGBColor(0, 79, 159) # Skyworth Blue
    
    cover_p3 = doc.add_paragraph()
    cover_p3.paragraph_format.space_after = Pt(110)
    cover_p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = cover_p3.add_run("Guía Integral de Arquitectura, Diccionario de Datos, Flujos Operativos y Espacios para Evidencias Visuales")
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Proyecto:", "Sistema de Bono e Incentivos Vendedores Skyworth"),
        ("Tipo de Sistema:", "Plataforma Web / PWA Móvil + API RESTful"),
        ("Backend:", "Java 17 / Spring Boot 3.3.6 / Spring Security"),
        ("Frontend:", "React 18 / TypeScript / Vite / Tailwind CSS / shadcn/ui"),
        ("Base de Datos:", "PostgreSQL 14+ (Pool HikariCP)"),
        ("Versión del Software:", "v1.0.0 (Producción / Release)"),
        ("Fecha de Edición:", "Septiembre 2026"),
        ("Entorno de Producción:", "https://bono.skyworth.bo / https://bono.contabol.com")
    ]
    style_table(meta_table, [2.2, 4.3], ["Atributo", "Especificación Técnica"], meta_data, header_bg="0F172A", alt_bg="F1F5F9")
    
    doc.add_page_break()
    
    # ================= 1. RESUMEN EJECUTIVO =================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Resumen Ejecutivo y Objetivos del Sistema")
    r.font.color.rgb = RGBColor(0, 79, 159)
    
    p = doc.add_paragraph(
        "El Sistema de Bono Vendedores Skyworth es una plataforma digital de misión crítica implementada para "
        "gestionar de punta a punta el incentivo económico y de fidelización a los promotores comerciales que venden "
        "televisores y electrodomésticos Skyworth en puntos de venta autorizados a nivel nacional."
    )
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    
    p2 = doc.add_paragraph(
        "Su principal virtud radica en transformar un proceso tradicionalmente disperso (fotografías por mensajería, "
        "planillas Excel manuales y retrasos prolongados en pagos) en un flujo automatizado con validación en tiempo real "
        "de números de serie y evidencias fotográficas, blindaje contra seriales duplicados o robados, asignación dinámica "
        "de bonos según campañas comerciales activas, y liquidaciones bancarias consolidadas."
    )
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(10)
    
    create_callout(
        doc,
        "La solución opera bajo un esquema desacoplado con frontend reactivo en React + Vite y backend robusto en Spring Boot 3, "
        "diseñado bajo los principios Mobile-First para facilitar la captura ágil en los salones de venta.",
        "ARQUITECTURA DE MISIÓN CRÍTICA"
    )
    
    # ================= 2. DESGLOSE DETALLADO DE MÓDULOS =================
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. Desglose Detallado de Módulos del Sistema")
    r.font.color.rgb = RGBColor(0, 79, 159)
    
    doc.add_paragraph(
        "A continuación se detalla la funcionalidad de cada uno de los módulos que integran el ecosistema. "
        "En cada sección se incluye un espacio reservado para insertar las capturas de pantalla de la aplicación."
    )
    
    # --- MODULO 2.1: AUTENTICACIÓN Y SEGURIDAD ---
    doc.add_heading("2.1. Módulo de Autenticación, Seguridad y Recuperación de Contraseña", level=2)
    p_mod1 = doc.add_paragraph(
        "Este módulo controla la barrera perimetral de acceso al sistema mediante tokens criptográficos JWT (JSON Web Tokens). "
        "Incluye control de sesiones expiradas, renovación de credenciales, y un mecanismo seguro de recuperación de contraseña "
        "que genera un token temporal de un solo uso con caducidad de 60 minutos, despachado por correo electrónico."
    )
    p_mod1.paragraph_format.line_spacing = 1.15
    
    p_mod1_f = doc.add_paragraph()
    p_mod1_f.add_run("• Funcionalidades clave: ").bold = True
    p_mod1_f.add_run(
        "Login unificado para todos los perfiles, validación de contraseñas con hash BCrypt, protección contra fuerza bruta, "
        "y enlace de autoservicio '¿Olvidaste tu contraseña?' con formulario de reseteo seguro."
    )
    
    create_image_placeholder(
        doc,
        "Pantalla de Inicio de Sesión (Login) y Recuperación de Contraseña",
        "Capture la pantalla inicial con campos de correo y contraseña, y la vista de recuperación."
    )
    
    # --- MODULO 2.2: ONBOARDING Y REGISTRO DE VENDEDORES ---
    doc.add_heading("2.2. Módulo de Registro y Solicitudes de Aprobación de Vendedores", level=2)
    p_mod2 = doc.add_paragraph(
        "Permite que nuevos promotores comerciales soliciten su afiliación a la plataforma Skyworth. "
        "El vendedor completa su información personal básica: Cédula de Identidad (CI), número de teléfono/WhatsApp, correo, "
        "tienda comercial asignada, ciudad y talla de indumentaria corporativa (polera oficial Skyworth)."
    )
    p_mod2.paragraph_format.line_spacing = 1.15
    
    p_mod2_f = doc.add_paragraph()
    p_mod2_f.add_run("• Flujo de Aprobación: ").bold = True
    p_mod2_f.add_run(
        "Al registrarse, la cuenta queda en estado 'pendingApproval = true'. El Administrador visualiza las solicitudes en su panel "
        "('RegistrationRequestsPage'), verifica la pertenencia del vendedor al canal de distribución y autoriza o deniega el acceso con un solo clic."
    )
    
    create_image_placeholder(
        doc,
        "Formulario de Registro de Nuevo Vendedor y Panel Administrativo de Aprobación",
        "Capture la vista pública de registro y la tabla administrativa de solicitudes pendientes."
    )

    # --- MODULO 2.3: PORTAL DEL VENDEDOR ---
    doc.add_heading("2.3. Módulo del Portal del Vendedor (Dashboard, Mis Ventas y Mis Pagos)", level=2)
    p_mod3 = doc.add_paragraph(
        "Es el centro de control móvil para los vendedores en tienda. Diseñado con una interfaz táctil limpia, ofrece tres pantallas fundamentales:"
    )
    p_mod3.paragraph_format.line_spacing = 1.15
    
    p_mod3_sub1 = doc.add_paragraph()
    p_mod3_sub1.add_run("1. Dashboard Principal: ").bold = True
    p_mod3_sub1.add_run("Resumen en tiempo real del saldo total de bonos acumulados en Bolivianos (Bs.), puntos vigentes, total de ventas registradas y atajos rápidos.")
    
    p_mod3_sub2 = doc.add_paragraph()
    p_mod3_sub2.add_run("2. Mis Ventas: ").bold = True
    p_mod3_sub2.add_run("Listado paginado con pestañas por estado: Pendientes (en revisión), Aprobadas (listas para pago), Rechazadas (con el motivo explicado) y Cerradas (pagadas).")
    
    p_mod3_sub3 = doc.add_paragraph()
    p_mod3_sub3.add_run("3. Mis Pagos: ").bold = True
    p_mod3_sub3.add_run("Historial de liquidaciones bancarias recibidas, fechas de pago, montos exactos y visualización del comprobante bancario emitido por Skyworth.")
    
    create_image_placeholder(
        doc,
        "Dashboard del Vendedor y Sección 'Mis Ventas'",
        "Capture la vista del vendedor mostrando el saldo acumulado en Bs., tarjetas de resumen y tabla de ventas."
    )

    doc.add_page_break()

    # --- MODULO 2.4: REGISTRO DE VENTAS Y EVIDENCIAS MULTIMEDIA ---
    doc.add_heading("2.4. Módulo de Registro de Ventas, Validación de Serial y Carga Multimedia", level=2)
    p_mod4 = doc.add_paragraph(
        "Es el núcleo transaccional del sistema para los promotores comerciales. Permite registrar una venta realizada en tienda "
        "mediante la digitación y validación en tiempo real del número de serie, garantizando la legitimidad del televisor y adjuntando "
        "las pruebas fotográficas de la transacción."
    )
    p_mod4.paragraph_format.line_spacing = 1.15
    
    p_mod4_steps = doc.add_paragraph()
    p_mod4_steps.add_run("• Ingreso y Validación de Serial en Tiempo Real: ").bold = True
    p_mod4_steps.add_run(
        "El vendedor ingresa el número de serie en el campo de búsqueda y presiona el botón 'Validar'. El sistema consulta de inmediato "
        "al backend (/serials/{serial}/validate), verificando si el serial es auténtico, autodetectando el modelo exacto del televisor "
        "(ej. Q6600H, G22, etc.) y mostrando las campañas activas con los puntos y bonos en Bs. que le corresponden.\n"
    )
    p_mod4_steps.add_run("• Carga de 3 Evidencias Fotográficas Obligatorias: ").bold = True
    p_mod4_steps.add_run(
        "(1) Foto de la etiqueta trasera del televisor (TAG con serial legible), (2) Foto de la póliza de garantía sellada y fechada por la tienda, "
        "y (3) Foto de la nota de venta o factura al cliente final.\n"
    )
    p_mod4_steps.add_run("• Optimización y Compresión en el Dispositivo Móvil: ").bold = True
    p_mod4_steps.add_run(
        "Convierte imágenes de iPhone en formato HEIC a JPEG mediante 'heic2any' y redimensiona en el cliente a un máximo de 1600px "
        "con 75% de calidad (mediante Canvas HTML5), reduciendo el peso de los archivos antes de enviarlos para ahorrar datos móviles."
    )
    
    create_image_placeholder(
        doc,
        "Pantalla de Registro de Venta con Validación de Serial y Carga de las 3 Fotografías",
        "Capture la interfaz donde el vendedor ingresa el serial, valida el modelo/bono proyectado y selecciona las 3 evidencias fotográficas."
    )

    # --- MODULO 2.5: GESTIÓN DE CAMPAÑAS Y MOTOR DE BONOS ---
    doc.add_heading("2.5. Módulo de Campañas Comerciales y Motor de Reglas de Bonificación", level=2)
    p_mod5 = doc.add_paragraph(
        "Permite a la gerencia comercial configurar dinámicamente las temporadas de bonificación (ej. 'Campaña Navidad', 'Promo Día de la Madre'). "
        "Cada campaña establece una fecha de inicio, una fecha de finalización y una tabla específica de productos participantes, "
        "donde a cada modelo de televisor se le asigna un monto específico de Bono en Bs. y Puntos de recompensa."
    )
    p_mod5.paragraph_format.line_spacing = 1.15
    
    p_mod5_f = doc.add_paragraph()
    p_mod5_f.add_run("• Cálculo Automatizado: ").bold = True
    p_mod5_f.add_run(
        "Cuando el vendedor registra una venta, el motor evalúa la fecha de venta y el modelo para cruzarlo con las campañas vigentes. "
        "El vendedor visualiza de inmediato cuánto dinero y cuántos puntos ganará por esa venta específica."
    )
    
    create_image_placeholder(
        doc,
        "Panel Administrativo de Configuración de Campañas Comerciales y Asignación de Bonos",
        "Capture la lista de campañas y el modal de configuración de productos participantes con sus importes en Bs."
    )

    # --- MODULO 2.6: MESA DE REVISIÓN TERRITORIAL ---
    doc.add_heading("2.6. Módulo de Mesa de Revisión y Validación Territorial (Reviewers)", level=2)
    p_mod6 = doc.add_paragraph(
        "Herramienta especializada para el equipo de revisores (Reviewers) y administradores encargados del control de calidad. "
        "Los revisores están asignados por ciudades o grupos de ciudades (ej. La Paz, Santa Cruz, Cochabamba), lo que segmenta automáticamente "
        "las ventas que les corresponde fiscalizar."
    )
    p_mod6.paragraph_format.line_spacing = 1.15
    
    p_mod6_f = doc.add_paragraph()
    p_mod6_f.add_run("• Visor de Fotografías de Alta Definición: ").bold = True
    p_mod6_f.add_run("Permite hacer zoom profundo (+/-), rotación 90° e inspección detallada para cotejar que el serial de la etiqueta coincida exactamente con el de la póliza sellada.\n")
    p_mod6_f.add_run("• Aprobación o Rechazo Asistido: ").bold = True
    p_mod6_f.add_run("Con un clic se aprueba la venta o se rechaza seleccionando un motivo predefinido ('Foto borrosa', 'Falta sello en póliza', 'Serial no coincide') o nota personalizada, la cual se archiva en el historial de rechazos.")
    
    create_image_placeholder(
        doc,
        "Mesa de Revisión (Reviews) con Visor de Fotos en Alta Resolución y Opciones de Aprobación",
        "Capture la pantalla donde el revisor inspecciona las fotos con zoom y aprueba o rechaza la venta."
    )

    doc.add_page_break()

    # --- MODULO 2.7: LIQUIDACIÓN Y GESTIÓN DE PAGOS ---
    doc.add_heading("2.7. Módulo de Gestión de Pagos, Lotes de Liquidación y Comprobantes Bancarios", level=2)
    p_mod7 = doc.add_paragraph(
        "Controla el ciclo financiero de pago de comisiones. Las ventas aprobadas por los revisores se consolidan en el módulo de pagos, "
        "donde el departamento contable puede generar cortes de liquidación por período, ciudad o vendedor."
    )
    p_mod7.paragraph_format.line_spacing = 1.15
    
    p_mod7_f = doc.add_paragraph()
    p_mod7_f.add_run("• Flujo Contable: ").bold = True
    p_mod7_f.add_run(
        "El sistema exporta la nómina consolidada con número de cuenta bancaria, CI y monto total a abonar. "
        "Una vez ejecutada la transferencia bancaria, el administrador sube el comprobante de respaldo (PDF o imagen) y liquida el lote, "
        "transicionando todas las ventas asociadas al estado irrevocable 'CERRADA'."
    )
    
    create_image_placeholder(
        doc,
        "Módulo de Liquidación y Gestión de Pagos a Vendedores",
        "Capture la pantalla de administración de pagos con montos consolidados y carga del comprobante de transferencia."
    )

    # --- MODULO 2.8: CATÁLOGO DE PRODUCTOS Y SERIALES LEGÍTIMOS ---
    doc.add_heading("2.8. Módulo Maestro de Productos y Padrón Oficial de Seriales", level=2)
    p_mod8 = doc.add_paragraph(
        "Administra el catálogo maestro de modelos Skyworth (Smart TVs de 32\" a 86\", tecnologías QLED, UHD 4K, Google TV) "
        "y el padrón oficial de números de serie legítimos importados al país por la marca."
    )
    p_mod8.paragraph_format.line_spacing = 1.15
    
    p_mod8_f = doc.add_paragraph()
    p_mod8_f.add_run("• Carga Masiva Excel: ").bold = True
    p_mod8_f.add_run(
        "Mediante Apache POI, el administrador puede importar en segundos archivos Excel con miles de números de serie despachados "
        "desde fábrica o bodega central, asociándolos a su modelo respectivo para nutrir el padrón de validación previa."
    )
    
    create_image_placeholder(
        doc,
        "Catálogo de Productos y Vista de Importación de Seriales Oficiales",
        "Capture la tabla de modelos de productos y la herramienta de carga masiva de seriales desde Excel."
    )

    # --- MODULO 2.9: ESCUDO ANTI-FRAUDE ---
    doc.add_heading("2.9. Módulo de Seriales Restringidos (Escudo Anti-Fraude)", level=2)
    p_mod9 = doc.add_paragraph(
        "Constituye la barrera de defensa contra fraudes, robos o mercadería siniestrada. Permite registrar números de serie en una "
        "'Lista Negra' permanente o temporal, indicando la causa (ej. 'Mercadería con daño de transporte', 'Producto reportado robado', 'Baja contable')."
    )
    p_mod9.paragraph_format.line_spacing = 1.15
    
    p_mod9_f = doc.add_paragraph()
    p_mod9_f.add_run("• Bloqueo Preventivo: ").bold = True
    p_mod9_f.add_run(
        "Si cualquier usuario intenta registrar un televisor cuyo serial figure en esta tabla, el sistema rechaza la transacción de inmediato "
        "y genera un registro de auditoría con la IP, usuario y hora exacta del intento."
    )
    
    create_image_placeholder(
        doc,
        "Lista Negra de Seriales Restringidos y Motivos de Bloqueo",
        "Capture el panel de seriales restringidos con las opciones para bloquear o desbloquear equipos."
    )

    doc.add_page_break()

    # --- MODULO 2.10: SERVICIO TÉCNICO Y EQUIPOS DAÑADOS ---
    doc.add_heading("2.10. Módulo de Soporte Técnico y Control de Equipos Dañados (Escáner de Código de Barras)", level=2)
    p_mod10 = doc.add_paragraph(
        "Diseñado para coordinar las garantías y reclamos de clientes con fallas de origen. Los técnicos autorizados disponen de un portal "
        "exclusivo para registrar televisores averiados devueltos por clientes o tiendas comerciales. En este módulo reside la tecnología "
        "de escaneo automático de código de barras."
    )
    p_mod10.paragraph_format.line_spacing = 1.15
    
    p_mod10_f = doc.add_paragraph()
    p_mod10_f.add_run("• Escáner de Código de Barras con Cámara y Reconocimiento de Imagen (ZXing): ").bold = True
    p_mod10_f.add_run(
        "Este módulo incorpora el componente 'BarcodeTagScanner' con la librería ZXing, permitiendo al técnico enfocar con la cámara del dispositivo "
        "el código de barras de la etiqueta trasera del televisor para decodificar automáticamente el serial. Además, el backend "
        "(EquipoDanadoService) cuenta con un decodificador avanzado con rotaciones en 0°, 90°, 180° y 270° que extrae el serial directamente "
        "de la fotografía del TAG subida.\n"
    )
    p_mod10_f.add_run("• Catálogo de Fallas y Diagnóstico: ").bold = True
    p_mod10_f.add_run(
        "Estandariza los tipos de desperfectos (falla de pantalla, backlight, audio, tarjeta madre o software). "
        "Permite adjuntar fotografías del defecto físico, registrar el serial del equipo y asentar si el equipo procede a cambio o reparación."
    )
    
    create_image_placeholder(
        doc,
        "Módulo de Servicio Técnico con Escáner de Código de Barras y Registro de Fallas",
        "Capture la pantalla del componente BarcodeTagScanner y el formulario de diagnóstico de equipos averiados."
    )

    # --- MODULO 2.11: GAMIFICACIÓN Y RANKING ---
    doc.add_heading("2.11. Módulo de Gamificación, Leaderboard y Ranking de Vendedores", level=2)
    p_mod11 = doc.add_paragraph(
        "Elemento motivacional clave que fomenta la sana competencia entre los vendedores de las diferentes ciudades y cadenas comerciales. "
        "Calcula la tabla de posiciones en tiempo real en función de las ventas aprobadas y los puntos acumulados."
    )
    p_mod11.paragraph_format.line_spacing = 1.15
    
    p_mod11_f = doc.add_paragraph()
    p_mod11_f.add_run("• Métricas Visibles: ").bold = True
    p_mod11_f.add_run(
        "Podio de honor (Top 3 vendedores), ranking por ciudad, medallas virtuales y visualización del avance hacia metas mensuales "
        "establecidas por la dirección comercial."
    )
    
    create_image_placeholder(
        doc,
        "Leaderboard y Ranking de Posiciones de Vendedores Skyworth",
        "Capture la vista del ranking con los mejores vendedores, puntos y posiciones destacadas."
    )

    # --- MODULO 2.12: DASHBOARD ADMINISTRATIVO, MÉTRICAS Y AUDITORÍA ---
    doc.add_heading("2.12. Módulo de Métricas, Reportes y Bitácora de Auditoría", level=2)
    p_mod12 = doc.add_paragraph(
        "Panel central de Business Intelligence para directores y supervisores. Cuenta con dashboards interactivos impulsados por Recharts "
        "que muestran el volumen de ventas por ciudad, los modelos de mayor rotación, y los montos totales de bonos devengados vs. pagados."
    )
    p_mod12.paragraph_format.line_spacing = 1.15
    
    p_mod12_f = doc.add_paragraph()
    p_mod12_f.add_run("• Auditoría Inmutable: ").bold = True
    p_mod12_f.add_run(
        "La tabla 'auditoria' registra cada cambio crítico (quién aprobó una venta, quién modificó un bono, quién bloqueó un serial) "
        "garantizando completa transparencia. Incluye botón para exportar reportes ejecutivos consolidados a Microsoft Excel."
    )
    
    create_image_placeholder(
        doc,
        "Dashboard Ejecutivo con Gráficos KPI y Bitácora de Auditoría",
        "Capture los gráficos analíticos de ventas por región y la tabla de auditoría con eventos registrados."
    )

    # --- MODULO 2.13: CONFIGURACIÓN GLOBAL Y POPUPS ---
    doc.add_heading("2.13. Módulo de Configuración Global, Popups y Notificaciones", level=2)
    p_mod13 = doc.add_paragraph(
        "Permite ajustar parámetros de operación del sistema sin necesidad de modificar código fuente ni reiniciar el servidor."
    )
    p_mod13.paragraph_format.line_spacing = 1.15
    
    p_mod13_f = doc.add_paragraph()
    p_mod13_f.add_run("• Herramientas de Configuración: ").bold = True
    p_mod13_f.add_run(
        "Publicación de ventanas emergentes (Popups) que ven los vendedores al abrir la app con avisos de nuevas campañas; "
        "gestión de correos destinatarios de notificaciones automáticas; y límite máximo de semanas permitidas para registrar ventas retroactivas."
    )
    
    create_image_placeholder(
        doc,
        "Panel de Configuración Global y Gestión de Popups Informativos",
        "Capture la pantalla de configuración de parámetros globales y creación de anuncios popup para vendedores."
    )

    doc.add_page_break()

    # ================= 3. GUÍA PARA INSERTAR IMÁGENES =================
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. Guía Paso a Paso para Insertar las Imágenes en Word")
    r.font.color.rgb = RGBColor(0, 79, 159)
    
    p_guide = doc.add_paragraph(
        "Los recuadros grises punteados incluidos en esta documentación han sido configurados con ancho y márgenes idóneos "
        "para recibir las capturas de pantalla de la aplicación. Para colocar sus imágenes siga uno de estos dos sencillos métodos:"
    )
    p_guide.paragraph_format.line_spacing = 1.15
    
    p_m1 = doc.add_paragraph()
    p_m1.add_run("• Método 1: Pegado Directo con el Portapapeles (Más Rápido):").bold = True
    p_m1.add_run(
        "\n  1. Tome una captura de pantalla en su navegador (presionando Windows + Shift + S o Print Screen)."
        "\n  2. Abra este documento en Microsoft Word."
        "\n  3. Haga clic dentro del recuadro punteado correspondiente."
        "\n  4. Presione Ctrl + V para pegar la imagen. La imagen se ajustará automáticamente al ancho del documento."
    )
    
    p_m2 = doc.add_paragraph()
    p_m2.add_run("• Método 2: Menú Insertar Imagen de Word:").bold = True
    p_m2.add_run(
        "\n  1. Guarde la captura de pantalla en su computadora (en formato PNG o JPG)."
        "\n  2. Coloque el cursor dentro del recuadro de la figura correspondiente."
        "\n  3. Vaya a la barra superior de Word: Insertar > Imágenes > Este dispositivo..."
        "\n  4. Seleccione el archivo guardado y haga clic en Insertar."
    )

    doc.add_page_break()

    # ================= 4. MATRIZ EJECUTIVA SINTETIZADA DE APIS REST =================
    h4_api = doc.add_heading(level=1)
    r_api = h4_api.add_run("4. Catálogo Sintetizado de APIs REST (Matriz Ejecutiva)")
    r_api.font.color.rgb = RGBColor(0, 79, 159)
    
    doc.add_paragraph(
        "A fin de proporcionar una lectura ágil y de alto nivel, los 123 endpoints del backend se han sintetizado en "
        "la siguiente Matriz Ejecutiva de Recursos. Cada fila agrupa las operaciones CRUD y acciones especializadas "
        "según su controlador y dominio funcional:"
    )

    tbl_matrix = doc.add_table(rows=0, cols=5)
    matrix_data = [
        ("Autenticación y Cuentas", "/api/auth", "POST, PUT", "Público / Autenticado", "Login con JWT, registro de vendedores, reseteo de clave por email (60 min) y cambio de contraseña."),
        ("Usuarios y Roles (RBAC)", "/api/users", "GET, POST, DELETE", "ADMIN", "Administración de usuarios internos (Admin, Reviewer, Supervisor, Técnico), asignación de roles y bajas."),
        ("Gestión de Vendedores", "/api/vendors", "GET, POST, PUT, PATCH", "ADMIN, REVIEWER", "Padrón de promotores, aprobación/rechazo de registro, activación, edición de perfil y estadísticas por ciudad."),
        ("Portal del Vendedor", "/api/vendor", "GET, PUT", "VENDOR", "Perfil del vendedor en sesión (/me), saldo acumulado de bonos en Bs., puntos (/stats) e historial de ventas (/sales)."),
        ("Registro y Control de Ventas", "/api/sales", "GET, POST, PUT, DELETE", "VENDOR, REVIEWER, ADMIN", "Alta de ventas con fotos (Multipart), mesa de revisión (aprobación/rechazo justificado), reversión y exportación a Excel."),
        ("Padrón de Seriales", "/api/serials", "GET, POST, PUT, DELETE", "ADMIN, Autenticado", "Validación en tiempo real (/validate), importación masiva en Excel (Apache POI), estadísticas y exportes."),
        ("Escudo Anti-Fraude", "/api/restricted-serials", "GET, POST, DELETE", "ADMIN, REVIEWER", "Lista negra de seriales bloqueados (por robo, daño o fraude), bloqueo masivo y desbloqueo."),
        ("Campañas Comerciales", "/api/campaigns", "GET, POST, PUT, DELETE", "ADMIN, Autenticado", "Configuración de temporadas de incentivos, vigencias, y asignación de bonos en Bs. y puntos por modelo de televisor."),
        ("Pagos y Liquidaciones", "/api/payments", "GET, POST", "ADMIN, VENDOR", "Corte contable de bonos aprobados, pago bancario con subida de comprobante y exportación de nómina a Excel."),
        ("Catálogo de Productos", "/api/products, /product-types", "GET, POST, PUT, DELETE", "ADMIN, Autenticado", "Maestro de modelos de Smart TV Skyworth, pulgadas, familias tecnológicas (QLED/UHD) y estados de activación."),
        ("Tiendas y Zonificación", "/api/tiendas, /cities, /city-groups", "GET, POST, PUT, PATCH, DELETE", "ADMIN, REVIEWER", "Catálogo de tiendas autorizadas, padrón de ciudades y grupos territoriales para asignación a supervisores."),
        ("Servicio Técnico / Dañados", "/api/admin/equipos-danados", "GET, POST, PUT, DELETE", "TECNICO, ADMIN", "Registro de televisores averiados con fotos, catálogo de fallas, clientes, garantías y escáner ZXing de códigos de barra."),
        ("Gamificación y Reportes", "/api/ranking, /reports", "GET, POST", "ADMIN, SUPERVISOR, Autenticado", "Leaderboard de ventas mensual en vivo, disparo manual de reportes por correo y envíos automáticos semanales."),
        ("Popups y Notificaciones", "/api/admin/popups, /email-recipients", "GET, POST, PUT, DELETE", "ADMIN, Público", "Gestión de anuncios emergentes (banners informativos para vendedores) y lista de correos receptores de alertas."),
        ("Configuración y Auditoría", "/api/admin/config, /audit, /dashboard", "GET, PUT", "ADMIN, SUPERVISOR", "Parámetros globales (semanas límite de venta), matriz de permisos, métricas KPI y bitácora inmutable de eventos.")
    ]
    style_table(tbl_matrix, [1.4, 1.3, 0.9, 1.0, 2.4], ["Recurso / Módulo", "URI Base", "Métodos", "Rol Mínimo", "Alcance y Operaciones Clave"], matrix_data)

    create_callout(
        doc,
        "DOCUMENTACIÓN INTERACTIVA ABIERTA (SWAGGER / OPENAPI):\n"
        "Para inspeccionar la especificación técnica atómica de cada uno de los 123 endpoints (cuerpos JSON de petición/respuesta, "
        "códigos de error HTTP y parámetros de consulta), el backend cuenta con SpringDoc OpenAPI activo en tiempo de ejecución en:\n"
        "• Interfaz Gráfica Swagger UI: http://localhost:7001/api/swagger-ui.html (o en el dominio productivo /api/swagger-ui.html)\n"
        "• Especificación JSON OpenAPI: /api/v3/api-docs",
        "REFERENCIA TÉCNICA DETALLADA"
    )

    doc.add_page_break()

    # ================= 5. GUÍA DE DESPLIEGUE Y PUESTA EN MARCHA =================
    h5_deploy = doc.add_heading(level=1)
    r_deploy = h5_deploy.add_run("5. Guía de Despliegue y Puesta en Marcha")
    r_deploy.font.color.rgb = RGBColor(0, 79, 159)

    doc.add_heading("5.1. Requisitos de Infraestructura y Servidor", level=2)
    req_p = doc.add_paragraph()
    req_p.add_run("• Sistema Operativo: ").bold = True
    req_p.add_run("Linux Ubuntu Server 22.04 LTS (recomendado en producción) o Windows Server 2019/2022.\n")
    req_p.add_run("• Entorno de Ejecución Java: ").bold = True
    req_p.add_run("OpenJDK 17 LTS de 64 bits.\n")
    req_p.add_run("• Entorno Node.js: ").bold = True
    req_p.add_run("Node.js v18.x o v20.x LTS con gestor de paquetes npm 9+.\n")
    req_p.add_run("• Motor de Base de Datos: ").bold = True
    req_p.add_run("PostgreSQL 14, 15 o 16 con codificación de caracteres UTF-8.\n")
    req_p.add_run("• Recursos de Hardware Sugeridos: ").bold = True
    req_p.add_run("Mínimo 4 GB de memoria RAM (8 GB recomendado para procesamiento de imágenes y exportaciones Excel masivas) y 40 GB en disco SSD.")

    doc.add_heading("5.2. Despliegue del Backend (Spring Boot)", level=2)
    doc.add_paragraph(
        "1. Posicionarse en el directorio del backend:\n"
        "   cd backend_bono_vendedor\n"
        "2. Empaquetar el artefacto ejecutable JAR omitiendo pruebas unitarias:\n"
        "   mvn clean package -DskipTests\n"
        "3. El archivo resultante 'bono-0.0.1-SNAPSHOT.jar' se genera en la carpeta 'target/'.\n"
        "4. Iniciar el servicio en segundo plano o como servicio systemd en Linux:\n"
        "   java -jar target/bono-0.0.1-SNAPSHOT.jar --spring.profiles.active=prod\n"
        "   (El servidor iniciará en el puerto 7001 bajo el context-path '/api')."
    )

    doc.add_heading("5.3. Despliegue del Frontend (React + Vite)", level=2)
    doc.add_paragraph(
        "1. Posicionarse en el directorio del frontend:\n"
        "   cd frontend_bono_vendedor3\n"
        "2. Configurar el archivo '.env' con la URL del backend:\n"
        "   VITE_API_URL=https://bono.skyworth.bo/api\n"
        "3. Instalar dependencias de producción y compilar activos estáticos:\n"
        "   npm install\n"
        "   npm run build\n"
        "4. Los archivos generados en la carpeta 'dist/' deben publicarse en el servidor web (Nginx o Apache) "
        "configurando redirección SPA ('try_files $uri $uri/ /index.html') para permitir el enrutamiento de React."
    )

    # ================= 6. MANTENIMIENTO, COPIAS DE SEGURIDAD Y BITÁCORAS =================
    doc.add_page_break()
    h6_maint = doc.add_heading(level=1)
    r_maint = h6_maint.add_run("6. Mantenimiento, Copias de Seguridad y Bitácoras")
    r_maint.font.color.rgb = RGBColor(0, 79, 159)

    doc.add_heading("6.1. Política de Bitácora Rotativa (Logs de Sistema)", level=2)
    p_log_desc = doc.add_paragraph(
        "El backend implementa un mecanismo rotativo y comprimido en 'application.properties' "
        "para garantizar la trazabilidad de peticiones y prevenir el desbordamiento de almacenamiento en disco:"
    )
    p_log_desc.paragraph_format.line_spacing = 1.15
    p_log_items = doc.add_paragraph()
    p_log_items.add_run("• Archivo de log activo: ").bold = True
    p_log_items.add_run("logs/skyworth.log\n")
    p_log_items.add_run("• Umbral de rotación: ").bold = True
    p_log_items.add_run("10 MB por archivo. Al superar este tamaño, se empaqueta automáticamente.\n")
    p_log_items.add_run("• Formato de compresión: ").bold = True
    p_log_items.add_run("GZIP con patrón de fecha: 'logs/skyworth-YYYY-MM-DD.i.log.gz'.\n")
    p_log_items.add_run("• Retención histórica: ").bold = True
    p_log_items.add_run("30 días naturales. Los registros con más de un mes de antigüedad se depuran automáticamente.")

    doc.add_heading("6.2. Estrategia de Copias de Seguridad (Backups)", level=2)
    doc.add_paragraph(
        "Se recomienda programar mediante cron las siguientes dos tareas de respaldo diario:\n"
        "1. Respaldo de Base de Datos PostgreSQL:\n"
        "   pg_dump -U postgres -d db_bono_skyworth -F c -b -v -f /backups/db_skyworth_$(date +%Y%m%d).dump\n"
        "2. Respaldo de Archivos Multimedia (/uploads):\n"
        "   tar -czf /backups/uploads_skyworth_$(date +%Y%m%d).tar.gz backend_bono_vendedor/uploads/"
    )

    create_callout(
        doc,
        "La carpeta '/uploads' contiene todas las evidencias fotográficas de venta (etiquetas TAG, pólizas selladas y recibos). "
        "Un respaldo combinado de PostgreSQL + uploads garantiza la restauración completa e íntegra del sistema ante cualquier contingencia.",
        "INTEGRIDAD DEL NEGOCIO"
    )

    # ================= 7. RESOLUCIÓN DE PROBLEMAS FRECUENTES (TROUBLESHOOTING) =================
    h7_trouble = doc.add_heading(level=1)
    r_trouble = h7_trouble.add_run("7. Guía de Solución de Problemas Frecuentes (Troubleshooting)")
    r_trouble.font.color.rgb = RGBColor(0, 79, 159)

    tbl_trouble = doc.add_table(rows=0, cols=3)
    trouble_data = [
        ("Serial no reconocido o inválido", "El vendedor ingresa un serial y el sistema indica 'No válido'.", "Verificar que el serial esté cargado en el padrón oficial (/api/serials) y que no se encuentre en la lista de seriales restringidos ni registrado previamente."),
        ("Error 401 Unauthorized / Sesión caducada", "La aplicación redirige al vendedor al login al intentar registrar una venta.", "El token JWT ha expirado. El frontend redirige automáticamente al login por seguridad. Basta con que el vendedor inicie sesión nuevamente."),
        ("Imágenes de iPhone (formato HEIC)", "El vendedor toma fotos con un iPhone y la aplicación no previsualiza.", "El sistema ya cuenta con la librería 'heic2any' en el cliente que convierte fotos HEIC a JPEG automáticamente. Si la memoria del teléfono es baja, se recomienda configurar la cámara del iPhone en 'El más compatible'."),
        ("Agotamiento del pool de base de datos", "El sistema responde con demoras o errores 500 en horas pico.", "HikariCP está configurado con 10 conexiones simultáneas. Si la concurrencia supera los 500 usuarios en simultáneo, incrementar 'spring.datasource.hikari.maximum-pool-size=25' en application.properties."),
        ("Correos de notificación no se envían", "Los vendedores no reciben el email de restablecimiento de contraseña.", "Ejecutar una prueba técnica mediante el endpoint GET /api/admin/config/test-smtp para diagnosticar si las credenciales SMTP o el puerto del proveedor de correo están activos.")
    ]
    style_table(tbl_trouble, [1.8, 2.0, 2.7], ["Incidencia Reportada", "Causa Probable", "Acción de Solución Recomendada"], trouble_data)

    doc.add_page_break()

    # ================= 8. GLOSARIO Y CONTROL DE VERSIONES =================
    h8_close = doc.add_heading(level=1)
    r_close = h8_close.add_run("8. Glosario de Términos y Control de Versiones")
    r_close.font.color.rgb = RGBColor(0, 79, 159)

    doc.add_heading("8.1. Glosario Técnico y del Negocio", level=2)
    p_glo = doc.add_paragraph()
    p_glo.add_run("• TAG (Etiqueta Trasera): ").bold = True
    p_glo.add_run("Adhesivo de fábrica adherido al televisor donde figura el modelo, código de barras y número de serie único.\n")
    p_glo.add_run("• Póliza Sellada: ").bold = True
    p_glo.add_run("Documento de garantía oficial de Skyworth que debe llevar el sello y firma del punto de venta autorizado.\n")
    p_glo.add_run("• Venta CERRADA: ").bold = True
    p_glo.add_run("Estado final e irrevocable de una venta aprobada cuyo bono en Bolivianos ya fue transferido y liquidado con comprobante.\n")
    p_glo.add_run("• JWT (JSON Web Token): ").bold = True
    p_glo.add_run("Estándar de seguridad criptográfico que viaja en las cabeceras HTTP permitiendo autenticar peticiones sin sobrecargar la base de datos.\n")
    p_glo.add_run("• RBAC: ").bold = True
    p_glo.add_run("Control de acceso basado en roles (Role-Based Access Control) que segmenta permisos entre Admin, Revisor, Supervisor, Vendedor y Técnico.")

    doc.add_heading("8.2. Historial de Revisiones del Documento", level=2)
    tbl_hist = doc.add_table(rows=0, cols=4)
    hist_data = [
        ("v1.0.0", "Septiembre 2026", "Ninabit Solutions", "Emisión inicial de documentación técnica, arquitectura y manual de módulos."),
        ("v1.0.1", "Septiembre 2026", "Ninabit Solutions", "Incorporación de espacios para capturas de pantalla y precisión en módulos de escaneo."),
        ("v1.0.2", "Septiembre 2026", "Ninabit Solutions", "Sintetización ejecutiva de la matriz de APIs REST (123 endpoints) e inclusión de guías operativas.")
    ]
    style_table(tbl_hist, [1.0, 1.3, 1.7, 2.5], ["Versión", "Fecha", "Elaborado Por", "Resumen de Cambios"], hist_data)

    doc.add_heading("8.3. Aprobación y Aceptación Formal del Sistema", level=2)
    doc.add_paragraph(
        "El presente documento técnico certifica la arquitectura, especificaciones funcionales y operativas de la plataforma "
        "Sistema de Bono Vendedores Skyworth a su fecha de corte."
    )
    
    # Signatures table
    tbl_sign = doc.add_table(rows=0, cols=2)
    sign_data = [
        ("\n\n_____________________________________\nRESPONSABLE TÉCNICO DE DESARROLLO\nNinabit Solutions",
         "\n\n_____________________________________\nRESPONSABLE COMERCIAL / NEGOCIO\nSkyworth Bolivia")
    ]
    style_table(tbl_sign, [3.2, 3.3], ["Entregado Por", "Aprobado y Recibido Por"], sign_data, header_bg="0F172A", alt_bg="FFFFFF")

    doc.save(output_path)
    print(f"Document successfully created at {output_path}")

if __name__ == "__main__":
    target = r"c:\Users\pauln\OneDrive\Escritorio\ninabit\BONO SKYWORTH\backend_bono_vendedor\Documentacion_Tecnica_Skyworth_Bono_Ejecutiva.docx"
    build_document(target)
